"""Gera as fixtures anonimizadas da Área do Investidor (layout público conhecido).

    cd apps/api && uv run python tests/fixtures/b3/gerar.py

Os arquivos trazem de propósito o que um arquivo real traz e que **não pode vazar**:
nome, CPF, conta e instituição (todos falsos, abaixo). `test_importers_b3.py` prova que
nenhum desses valores aparece na prévia.

Quando houver arquivos reais para validar (antes da `v0.4.0`), confira os cabeçalhos
contra estes e ajuste os apelidos em `importers/b3/*.py`, não as fixtures.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).parent

NOME = "FULANA DE TESTE DA SILVA"
CPF = "123.456.789-09"
CONTA = "7654321"
CORRETORA = "CORRETORA FICTICIA CCTVM S.A."
CORRETORA_2 = "OUTRA CORRETORA DE TESTE S.A."


def _active(wb: Workbook, title: str) -> Worksheet:
    ws = wb.active
    assert isinstance(ws, Worksheet)
    ws.title = title
    return ws


def _preambulo(ws: Worksheet, titulo: str) -> None:
    ws.append([titulo])
    ws.append([f"Investidor: {NOME}"])
    ws.append([f"CPF: {CPF}"])
    ws.append([])


def posicao() -> None:
    wb = Workbook()
    acoes = _active(wb, "Acoes")
    _preambulo(acoes, "Posição consolidada em 19/09/2026")
    header = [
        "Produto",
        "Instituição",
        "Conta",
        "Código de Negociação",
        "CNPJ da Empresa",
        "Código ISIN / Distribuição",
        "Tipo",
        "Escriturador",
        "Quantidade",
        "Quantidade Disponível",
        "Quantidade Indisponível",
        "Motivo",
        "Preço de Fechamento",
        "Valor Atualizado",
    ]
    acoes.append(header)
    acoes.append(
        [
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            CORRETORA,
            CONTA,
            "PETR4",
            "33.000.167/0001-01",
            "BRPETRACNPR6",
            "PN",
            "BANCO BRADESCO",
            155,
            155,
            0,
            "",
            38.5,
            5967.5,
        ]
    )
    acoes.append(
        [
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            CORRETORA_2,
            "111",
            "PETR4",
            "33.000.167/0001-01",
            "BRPETRACNPR6",
            "PN",
            "BANCO BRADESCO",
            45,
            45,
            0,
            "",
            38.5,
            1732.5,
        ]
    )
    acoes.append(
        [
            "VALE3 - VALE S.A.",
            CORRETORA,
            CONTA,
            "VALE3",
            "33.592.510/0001-54",
            "BRVALEACNOR0",
            "ON",
            "BANCO BRADESCO",
            "300",
            "300",
            "0",
            "",
            "61,20",
            "18.360,00",
        ]
    )
    acoes.append(["", "", "", "", "", "", "", "", "", "", "", "", "Total", 26060.0])

    bdr = wb.create_sheet("BDR")
    bdr.append(header)
    bdr.append(
        [
            "AAPL34 - APPLE INC",
            CORRETORA,
            CONTA,
            "AAPL34",
            "",
            "BRAAPLBDR004",
            "DRN",
            "BANCO BRADESCO",
            20,
            20,
            0,
            "",
            70.1,
            1402.0,
        ]
    )

    etf = wb.create_sheet("ETF")
    etf.append(header)
    etf.append(
        [
            "BOVA11 - ISHARES BOVA",
            CORRETORA,
            CONTA,
            "BOVA11",
            "",
            "BRBOVACTF003",
            "CI",
            "BANCO BRADESCO",
            10,
            10,
            0,
            "",
            130.0,
            1300.0,
        ]
    )

    fundos = wb.create_sheet("Fundo de Investimento")
    fundos.append(
        [
            "Produto",
            "Instituição",
            "Conta",
            "Código de Negociação",
            "CNPJ do Fundo",
            "Código ISIN / Distribuição",
            "Tipo",
            "Administrador",
            "Quantidade",
            "Quantidade Disponível",
            "Quantidade Indisponível",
            "Motivo",
            "Preço de Fechamento",
            "Valor Atualizado",
        ]
    )
    fundos.append(
        [
            "HGLG11 - CSHG LOGISTICA FII",
            CORRETORA,
            CONTA,
            "HGLG11",
            "11.728.688/0001-47",
            "BRHGLGCTF004",
            "Cotas",
            "ADMINISTRADORA TESTE",
            30,
            30,
            0,
            "",
            160.0,
            4800.0,
        ]
    )

    tesouro = wb.create_sheet("Tesouro Direto")
    tesouro.append(
        [
            "Produto",
            "Instituição",
            "Código ISIN",
            "Indexador",
            "Vencimento",
            "Quantidade",
            "Quantidade Disponível",
            "Quantidade Indisponível",
            "Motivo",
            "Valor Aplicado",
            "Valor bruto",
            "Valor líquido",
            "Valor Atualizado",
        ]
    )
    tesouro.append(
        [
            "Tesouro IPCA+ 2035",
            CORRETORA,
            "BRSTNCNTB4U6",
            "IPCA",
            "15/05/2035",
            1.5,
            1.5,
            0,
            "",
            3000.0,
            3300.0,
            3250.0,
            3300.0,
        ]
    )
    tesouro.append(
        [
            "Tesouro Selic 2029",
            CORRETORA,
            "BRSTNCLF1RU6",
            "SELIC",
            "01/03/2029",
            0.3,
            0.3,
            0,
            "",
            4500.0,
            4700.0,
            4680.0,
            4700.0,
        ]
    )

    rf = wb.create_sheet("Renda Fixa")
    rf.append(
        [
            "Produto",
            "Instituição",
            "Emissor",
            "Código",
            "Indexador",
            "Tipo de regime",
            "Data de Emissão",
            "Vencimento",
            "Quantidade",
            "Quantidade Disponível",
            "Quantidade Indisponível",
            "Motivo",
            "Contraparte",
            "Preço Atualizado MTM",
            "Valor Atualizado MTM",
            "Preço Atualizado CURVA",
            "Valor Atualizado CURVA",
        ]
    )
    rf.append(
        [
            "CDB - BANCO EMISSOR TESTE S.A.",
            CORRETORA,
            "BANCO EMISSOR TESTE S.A.",
            "CDB924ABC12",
            "DI",
            "Depositado",
            "02/01/2025",
            "02/01/2029",
            5,
            5,
            0,
            "",
            "",
            1100.0,
            5500.0,
            1098.0,
            5490.0,
        ]
    )
    rf.append(
        [
            "LCI - BANCO VENCIDO S.A.",
            CORRETORA,
            "BANCO VENCIDO S.A.",
            "LCI000VENC1",
            "DI",
            "Depositado",
            "02/01/2022",
            "02/01/2024",
            1,
            1,
            0,
            "",
            "",
            "",
            "",
            1000.0,
            1000.0,
        ]
    )

    opcoes = wb.create_sheet("Opções")
    opcoes.append(["Produto", "Instituição", "Código de Negociação", "Quantidade"])
    opcoes.append(["PETRJ300", CORRETORA, "PETRJ300", 100])

    wb.save(HERE / "posicao.xlsx")


def negociacao() -> None:
    wb = Workbook()
    ws = _active(wb, "Negociação")
    _preambulo(ws, "Extrato de negociação")
    ws.append(
        [
            "Data do Negócio",
            "Tipo de Movimentação",
            "Mercado",
            "Prazo/Vencimento",
            "Instituição",
            "Código de Negociação",
            "Quantidade",
            "Preço",
            "Valor",
        ]
    )
    rows = [
        ("10/01/2024", "Compra", "Mercado à Vista", "-", "PETR4", 100, 30.0, 3000.0),
        # Duas execuções idênticas no mesmo dia: são dois negócios, não uma repetição.
        ("15/03/2024", "Compra", "Mercado à Vista", "-", "PETR4", 50, 34.0, 1700.0),
        ("15/03/2024", "Compra", "Mercado à Vista", "-", "PETR4", 50, 34.0, 1700.0),
        ("20/05/2024", "Compra", "Mercado Fracionário", "-", "PETR4F", 5, "R$ 36,40", 182.0),
        ("02/07/2024", "Venda", "Mercado à Vista", "-", "PETR4", 55, 38.0, 2090.0),
        ("05/08/2024", "Compra", "Mercado à Vista", "-", "AAPL34", 20, 55.3, 1106.0),
        ("06/08/2024", "Compra", "Mercado à Vista", "-", "BOVA11", 10, 120.0, 1200.0),
        ("07/08/2024", "Compra", "Opção de Compra", "18/10/2024", "PETRJ300", 100, 0.5, 50.0),
    ]
    for d, side, market, prazo, ticker, qty, price, value in rows:
        ws.append([d, side, market, prazo, CORRETORA, ticker, qty, price, value])
    wb.save(HERE / "negociacao.xlsx")


def movimentacao() -> None:
    wb = Workbook()
    ws = _active(wb, "Movimentação")
    ws.append(
        [
            "Entrada/Saída",
            "Data",
            "Movimentação",
            "Produto",
            "Instituição",
            "Quantidade",
            "Preço unitário",
            "Valor da Operação",
        ]
    )
    rows = [
        (
            "Credito",
            "20/03/2024",
            "Dividendo",
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            200,
            0.54,
            108.0,
        ),
        (
            "Credito",
            "20/03/2024",
            "Juros Sobre Capital Próprio",
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            200,
            0.2125,
            42.5,
        ),
        ("Credito", "15/04/2024", "Rendimento", "HGLG11 - CSHG LOGISTICA FII", 30, 1.1, 33.0),
        (
            "Credito",
            "15/03/2024",
            "Transferência - Liquidação",
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            50,
            34.0,
            1700.0,
        ),
        (
            "Debito",
            "02/07/2024",
            "Transferência - Liquidação",
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            55,
            38.0,
            2090.0,
        ),
        (
            "Credito",
            "22/04/2024",
            "Dividendo - Transferido",
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            10,
            0.54,
            5.4,
        ),
        ("Credito", "30/12/2099", "Dividendo", "VALE3 - VALE S.A.", 300, 2.0, 600.0),
        ("Credito", "15/09/2026", "Dividendo", "VALE3 - VALE S.A.", 300, 2.1, "-"),
    ]
    for direction, d, event, product, qty, unit, value in rows:
        ws.append([direction, d, event, product, CORRETORA, qty, unit, value])
    wb.save(HERE / "movimentacao.xlsx")


def proventos_recebidos() -> None:
    wb = Workbook()
    ws = _active(wb, "Proventos Recebidos")
    _preambulo(ws, "Proventos recebidos")
    ws.append(
        [
            "Produto",
            "Pagamento",
            "Tipo de Evento",
            "Instituição",
            "Quantidade",
            "Preço unitário",
            "Valor líquido",
        ]
    )
    # A mesma linha que está na movimentação: enviar os dois não pode duplicar.
    ws.append(
        [
            "PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS",
            datetime(2024, 3, 20),
            "Dividendo",
            CORRETORA,
            200,
            0.54,
            108.0,
        ]
    )
    wb.save(HERE / "proventos-recebidos.xlsx")


if __name__ == "__main__":
    posicao()
    negociacao()
    movimentacao()
    proventos_recebidos()
    print("fixtures geradas em", HERE)  # noqa: T201 - script de linha de comando
