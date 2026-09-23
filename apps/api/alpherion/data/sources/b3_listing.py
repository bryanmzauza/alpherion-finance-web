"""Cadastro dos papéis listados na B3: ações, units, ETFs, BDRs, FIIs e fiagros.

É o que dá ao `securities` o que a CVM não tem — **ticker**, classe do papel e a
classificação setorial da B3. Sem isso não existe `/acoes/PETR4` nem `/setores/[slug]`.

Três fontes, cada uma para o que só ela tem, casadas pelo **código do emissor** (`PETR`):

1. **Cadastro de instrumentos** (`InstrumentsConsolidatedFile`, arquivo público diário
   de arquivos.b3.com.br): um registro por ticker negociado, com a categoria da própria
   B3 (`SHARES`, `UNIT`, `BDR`, `ETF EQUITIES`, `FUNDS`…), ISIN, razão social e nível de
   governança. É daqui que sai **o que existe** e **de que classe é** — a classe não é
   adivinhada pelo final do ticker (o `11` é unit, ETF e FII ao mesmo tempo).
2. **Classificação setorial** (planilha de `GetDownloadIndustryClassification`, a mesma
   que a página "Classificação setorial" da B3 oferece): setor → subsetor → segmento de
   cada emissor.
3. **Companhias listadas** (`GetInitialCompanies`): CNPJ e código CVM do emissor — o
   elo com as demonstrações da CVM. Esse endpoint lista *emissores* (3,5 mil, incluindo
   os que não negociam), não tickers: foi ler dele o ticker que deixou o cadastro inteiro
   com códigos de emissor classificados como BDR.

Só (1) é obrigatória. Sem (2), os papéis entram sem setor; sem (3), sem CNPJ — e o job
registra a falta. Nada disso inventa dado.

ADR-017: só o cadastro vem daqui; preço é COTAHIST e depende da licença.
"""

from __future__ import annotations

import csv
import logging
import re
import tempfile
import unicodedata
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Final

import httpx
from openpyxl import load_workbook

from alpherion.data.sources import http as source_http
from alpherion.data.sources.b3_api import (
    LISTED_BASE,
    B3UnavailableError,
    build_url,
    fetch_json,
    fetch_pages,
    field,
)

logger = logging.getLogger(__name__)

COMPANIES_PATH: Final = "listedCompaniesProxy/CompanyCall/GetInitialCompanies"
CLASSIFICATION_PATH: Final = "listedCompaniesProxy/CompanyCall/GetDownloadIndustryClassification"

ARQUIVOS_BASE: Final = "https://arquivos.b3.com.br/api/download"
INSTRUMENTS_FILE: Final = "InstrumentsConsolidatedFile"

#: Limites de download (§7.5). O arquivo de instrumentos tem ~20 MB (todas as opções e
#: futuros vêm junto); a planilha setorial, ~20 kB.
INSTRUMENTS_MAX_BYTES: Final = 80 * 1024 * 1024
CLASSIFICATION_MAX_BYTES: Final = 5 * 1024 * 1024

#: Quantos dias recuar atrás de um arquivo "Final" (feriado emendado com fim de semana).
LOOKBACK_DAYS: Final = 7

#: Categoria da B3 → `securities.type`. Fora daqui (recibo, direito, warrant, FIDC…) não
#: entra: melhor uma classe a menos do que um papel na página errada.
CATEGORY_TYPES: Final[dict[str, str]] = {
    "SHARES": "stock",
    "UNIT": "unit",
    "BDR": "bdr",
    "ETF EQUITIES": "etf",
    "ETF FOREIGN INDEX": "etf",
}

#: `FUNDS` reúne FII, fiagro, FI-Infra, FIP e FIDC. Só os dois primeiros têm página, e o
#: que os separa é o nome registrado — que a regulação obriga a dizer o que o fundo é.
FIAGRO_MARKERS: Final = ("FIAGRO",)
FII_MARKERS: Final = ("IMOB", "FII", "FDO INV IMOB", "IMOBILIARIO")

TICKER_RE: Final = re.compile(r"^[A-Z0-9]{4}\d{1,2}[A-Z]?$")
ISSUER_RE: Final = re.compile(r"^[A-Z0-9]{4}$")


@dataclass(frozen=True, slots=True)
class ListedSecurity:
    """Um papel listado, pronto para `securities`."""

    ticker: str
    type: str
    company_name: str
    trade_name: str | None = None
    cnpj: str | None = None
    cvm_code: int | None = None
    isin: str | None = None
    sector: str | None = None
    subsector: str | None = None
    segment: str | None = None
    sector_slug: str | None = None
    listing_segment: str | None = None
    #: ETF: slug do índice replicado. BDR: quantos BDRs equivalem a uma ação lá fora.
    etf_index_slug: str | None = None
    bdr_ratio: str | None = None

    @property
    def is_fund(self) -> bool:
        return self.type in {"fii", "fiagro"}


@dataclass(frozen=True, slots=True)
class Instrument:
    """Uma linha do cadastro de instrumentos, já filtrada para o mercado à vista."""

    ticker: str
    issuer: str
    category: str
    isin: str | None
    name: str
    governance: str | None


@dataclass(frozen=True, slots=True)
class Classification:
    """Onde a B3 classifica um emissor."""

    issuer: str
    sector: str
    subsector: str
    segment: str
    trade_name: str | None
    listing_segment: str | None


@dataclass(frozen=True, slots=True)
class Issuer:
    """Emissor listado: o elo com a CVM (`cvm_code`, CNPJ)."""

    code: str
    company_name: str
    trade_name: str | None
    cnpj: str | None
    cvm_code: int | None


# --- texto ------------------------------------------------------------------


def slugify(*parts: str | None) -> str | None:
    """`Petróleo, Gás e Biocombustíveis` → `petroleo-gas-e-biocombustiveis`."""
    text = " ".join(part for part in parts if part)
    if not text.strip():
        return None
    normalized = unicodedata.normalize("NFD", text)
    ascii_only = "".join(c for c in normalized if unicodedata.category(c) != "Mn")
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_only.lower()).strip("-")
    return slug[:80] or None


def _clean(value: Any) -> str | None:
    text = " ".join(str(value).split()) if value is not None else ""
    return text or None


def _digits(value: Any) -> str | None:
    only = "".join(c for c in str(value or "") if c.isdigit())
    return only or None


# --- cadastro de instrumentos -----------------------------------------------


def fund_type(name: str) -> str | None:
    """FII, fiagro ou nenhum dos dois, pelo nome registrado do fundo."""
    upper = unicodedata.normalize("NFD", name.upper())
    upper = "".join(c for c in upper if unicodedata.category(c) != "Mn")
    if any(marker in upper for marker in FIAGRO_MARKERS):
        return "fiagro"
    if any(marker in upper for marker in FII_MARKERS):
        return "fii"
    return None


def type_of(instrument: Instrument) -> str | None:
    """Classe do papel pela categoria da B3 — nunca pelo sufixo do ticker."""
    if instrument.category == "FUNDS":
        return fund_type(instrument.name)
    return CATEGORY_TYPES.get(instrument.category)


def parse_instruments(lines: Iterable[str]) -> tuple[str | None, list[Instrument]]:
    """Lê o arquivo de instrumentos: (status do arquivo, papéis do mercado à vista).

    A primeira linha é o status ("Status do Arquivo: Final" ou "Parcial"); a segunda, o
    cabeçalho. Fica só o **lote padrão** do mercado à vista (`SgmtNm = CASH`): o
    fracionário (`PETR4F`) e o lote de bloco são o mesmo papel com outro código.
    """
    iterator = iter(lines)
    first = next(iterator, "")
    status = first.split(":", 1)[1].strip() if ":" in first else None
    reader = csv.DictReader(iterator, delimiter=";")
    instruments: list[Instrument] = []
    for row in reader:
        if (row.get("SgmtNm") or "").strip() != "CASH":
            continue
        if (row.get("MktNm") or "").strip() != "EQUITY-CASH":
            continue
        ticker = (row.get("TckrSymb") or "").strip().upper()
        issuer = (row.get("Asst") or "").strip().upper()
        name = _clean(row.get("CrpnNm"))
        if not TICKER_RE.match(ticker) or not issuer or not name:
            continue
        instruments.append(
            Instrument(
                ticker=ticker,
                issuer=issuer,
                category=(row.get("SctyCtgyNm") or "").strip().upper(),
                isin=_clean(row.get("ISIN")),
                name=name,
                governance=_clean(row.get("CorpGovnLvlNm")),
            )
        )
    return status, instruments


def instruments_request_url(day: date) -> str:
    return (
        f"{ARQUIVOS_BASE}/requestname?fileName={INSTRUMENTS_FILE}"
        f"&date={day:%Y-%m-%d}&recaptchaToken="
    )


def _download_token(payload: Any) -> str:
    redirect = payload.get("redirectUrl") if isinstance(payload, dict) else None
    if not isinstance(redirect, str) or "token=" not in redirect:
        raise B3UnavailableError("arquivos.b3.com.br: resposta sem token de download")
    return redirect.split("token=", 1)[1]


def fetch_instruments(
    today: date, *, http: httpx.Client | None = None
) -> tuple[date, list[Instrument]]:
    """O cadastro de instrumentos do último pregão com arquivo **final**.

    O arquivo do dia sai "Parcial" durante o pregão; dia sem pregão responde 400. Recua
    até achar um final; se na janela só houver parcial, usa o mais recente — é o mesmo
    cadastro, só não fechado.
    """
    partial: tuple[date, list[Instrument]] | None = None
    with tempfile.TemporaryDirectory(prefix="alpherion-b3-instr-") as tmp:
        for back in range(LOOKBACK_DAYS + 1):
            day = today - timedelta(days=back)
            if day.weekday() >= 5:
                continue
            try:
                token = _download_token(fetch_json(instruments_request_url(day), http=http))
            except B3UnavailableError as error:
                logger.info("instrumentos de %s indisponíveis: %s", day, error)
                continue
            path = Path(tmp) / f"instrumentos-{day:%Y%m%d}.csv"
            source_http.download(
                f"{ARQUIVOS_BASE}/?token={token}",
                path,
                max_bytes=INSTRUMENTS_MAX_BYTES,
                http=http,
            )
            with path.open(encoding="latin-1", newline="") as file:
                status, instruments = parse_instruments(file)
            path.unlink(missing_ok=True)
            if not instruments:
                continue
            if status == "Final":
                return day, instruments
            partial = partial or (day, instruments)
    if partial is not None:
        logger.warning(
            "sem arquivo final de instrumentos na janela; usando o parcial de %s", partial[0]
        )
        return partial
    raise B3UnavailableError(f"nenhum cadastro de instrumentos nos últimos {LOOKBACK_DAYS} dias")


# --- classificação setorial -------------------------------------------------


def parse_classification(rows: Iterable[tuple[Any, ...]]) -> dict[str, Classification]:
    """Planilha da classificação setorial → emissor → setor/subsetor/segmento.

    Setor, subsetor e segmento vêm **mesclados**: só a primeira linha do bloco tem o
    valor, e as seguintes herdam. O cabeçalho se repete a cada quebra de página e é
    pulado. Linhas sem código de emissor válido não viram nada.
    """
    result: dict[str, Classification] = {}
    sector = subsector = segment = None
    for row in rows:
        cells = [_clean(c) for c in (*row, *([None] * 7))][:7]
        _, raw_sector, raw_subsector, raw_segment, trade_name, code, listing = cells
        if raw_sector and raw_sector.upper().startswith("SETOR"):
            continue  # cabeçalho repetido
        if raw_sector:
            sector = raw_sector
        if raw_subsector:
            subsector = raw_subsector
        if raw_segment:
            segment = raw_segment
        issuer = (code or "").upper()
        if not ISSUER_RE.match(issuer) or not (sector and subsector and segment):
            continue
        result[issuer] = Classification(
            issuer=issuer,
            sector=sector,
            subsector=subsector,
            segment=segment,
            trade_name=trade_name,
            listing_segment=listing,
        )
    return result


def segment_slugs(classes: Iterable[Classification]) -> dict[tuple[str, str], str]:
    """(subsetor, segmento) → slug de `/setores/[slug]`.

    O slug é só o segmento (`bancos`) — curto e estável. Quando o mesmo nome de segmento
    aparece em mais de um subsetor ("Serviços Diversos"), o subsetor entra na frente
    para os dois não virarem a mesma página.
    """
    pairs = {(c.subsector, c.segment) for c in classes}
    repeated = Counter(segment for _, segment in pairs)
    slugs: dict[tuple[str, str], str] = {}
    for subsector, segment in pairs:
        slug = slugify(subsector, segment) if repeated[segment] > 1 else slugify(segment)
        if slug:
            slugs[(subsector, segment)] = slug
    return slugs


def fetch_classification(*, http: httpx.Client | None = None) -> dict[str, Classification]:
    url = build_url(LISTED_BASE, CLASSIFICATION_PATH, {"language": "pt-br"})
    with tempfile.TemporaryDirectory(prefix="alpherion-b3-setor-") as tmp:
        path = source_http.download(
            url, Path(tmp) / "classificacao.xlsx", max_bytes=CLASSIFICATION_MAX_BYTES, http=http
        )
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:  # zip corrompido, HTML no lugar do xlsx
            raise B3UnavailableError(f"classificação setorial ilegível: {error}") from error
        try:
            sheet = workbook.worksheets[0]
            classes = parse_classification(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    if not classes:
        raise B3UnavailableError("classificação setorial veio vazia (formato mudou?)")
    logger.info("B3 classificação setorial: %d emissores", len(classes))
    return classes


# --- emissores --------------------------------------------------------------


def parse_issuer(record: dict[str, Any]) -> Issuer | None:
    code = str(field(record, "issuingCompany", "code") or "").strip().upper()
    name = _clean(field(record, "companyName"))
    if not ISSUER_RE.match(code) or not name:
        return None
    cvm = _digits(field(record, "codeCVM", "cvmCode"))
    return Issuer(
        code=code,
        company_name=name[:200],
        trade_name=_clean(field(record, "tradingName")),
        cnpj=_digits(field(record, "cnpj")),
        cvm_code=int(cvm) if cvm else None,
    )


def fetch_issuers(*, http: httpx.Client | None = None) -> dict[str, Issuer]:
    records = fetch_pages(LISTED_BASE, COMPANIES_PATH, {"language": "pt-br"}, http=http)
    issuers = {i.code: i for r in records if (i := parse_issuer(r)) is not None}
    logger.info("B3 emissores: %d", len(issuers))
    return issuers


# --- montagem ---------------------------------------------------------------


def build_listing(
    instruments: Iterable[Instrument],
    classification: dict[str, Classification],
    issuers: dict[str, Issuer],
) -> Iterator[ListedSecurity]:
    """Junta as três fontes num `ListedSecurity` por ticker."""
    slugs = segment_slugs(classification.values())
    seen: set[str] = set()
    for instrument in instruments:
        kind = type_of(instrument)
        if kind is None or instrument.ticker in seen:
            continue
        seen.add(instrument.ticker)
        issuer = issuers.get(instrument.issuer)
        # CNPJ e código CVM do emissor valem para companhia; fundo e ETF têm CNPJ próprio,
        # que o emissor listado (a gestora) não é.
        company = issuer if issuer and kind not in {"fii", "fiagro", "etf"} else None
        # Fundo não tem classificação setorial da B3: o segmento dele vem do informe
        # mensal da CVM (`cvm_fii_reports`).
        klass = None if kind in {"fii", "fiagro"} else classification.get(instrument.issuer)
        yield ListedSecurity(
            ticker=instrument.ticker,
            type=kind,
            company_name=(company.company_name if company else instrument.name)[:200],
            trade_name=(klass.trade_name if klass else None)
            or (issuer.trade_name if issuer else None),
            cnpj=issuer.cnpj if issuer and kind not in {"fii", "fiagro", "etf"} else None,
            cvm_code=issuer.cvm_code if issuer and kind not in {"fii", "fiagro", "etf"} else None,
            isin=instrument.isin,
            sector=klass.sector if klass else None,
            subsector=klass.subsector if klass else None,
            segment=klass.segment if klass else None,
            sector_slug=slugs.get((klass.subsector, klass.segment)) if klass else None,
            listing_segment=(klass.listing_segment if klass else None) or instrument.governance,
        )
